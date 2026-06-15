import os, sys
from datetime import datetime
import shutil

import pandas as pd
import openpyxl
import xlsxwriter

from .helper import *
from .helper import Serialize, Deserialize
from .log import *


"""
Converter

"""

def XlsxToJson(input_path, output_path):
    """Convert localization xlsx → id→korean JSON.

    Read columns directly via openpyxl by header NAME rather than positional
    integer keys. The original pandas-based implementation gated on
    `0 in input_record_keys` which fails the moment the first column header
    becomes anything other than a plain int — which is exactly what happens
    when apply_diff_to_xlsx (release sync) writes through openpyxl and the
    JP source column ends up rendered as `Unnamed: 0` by pandas. Every row
    silently dropped → empty JSON published.

    The current row shape produced by the release pipeline:
        col A  — JP source (header was '0' or an array formula; we match by
                 position-after-known-headers, not by name)
        col B  — '번역'  (Korean translation; export target)
        col C  — 'ID'    (lookup key)
        col D+ — unused

    Rows are skipped when:
      - ID cell isn't a non-empty string
      - 번역 cell is empty
      - 번역 carries the OBSOLETE marker via JP column (entry retired)
    """
    import openpyxl
    from .localization_release import OBSOLETE_MARKER

    wb = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        headers = [c.value for c in ws[1]]
        try:
            id_col = headers.index("ID")
            kr_col = headers.index("번역")
        except ValueError:
            LOG_ERROR(3, f"localization xlsx missing required headers (got: {headers})")
            raise

        # JP source column: anything that's not ID/번역 and lives at the
        # head of the sheet. Matches both legacy int(0) header and the
        # newer formula/Unnamed: 0 cases.
        jp_col = next(
            (i for i, h in enumerate(headers)
             if i != id_col and i != kr_col
             and (h == 0 or h == "0" or h is None or i == 0)),
            0,
        )

        data = {}
        skipped_no_id = skipped_no_trans = skipped_obsolete = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if id_col >= len(row) or kr_col >= len(row):
                continue
            key = row[id_col]
            if not isinstance(key, str) or not key:
                skipped_no_id += 1
                continue
            trans = row[kr_col]
            if not isinstance(trans, str) or not trans:
                skipped_no_trans += 1
                continue
            jp_val = row[jp_col] if jp_col < len(row) else ""
            if isinstance(jp_val, str) and jp_val.startswith(OBSOLETE_MARKER):
                skipped_obsolete += 1
                continue
            if trans.startswith("'"):
                data[key] = Deserialize(trans[1:])
            else:
                data[key] = Deserialize(trans)
    finally:
        wb.close()

    LOG_INFO(3, f"localization → JSON: wrote {len(data)} keys "
                f"(skipped: no-id={skipped_no_id}, no-trans={skipped_no_trans}, "
                f"obsolete={skipped_obsolete})")

    os.makedirs(os.path.split(output_path)[0], exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, allow_nan=False, indent=4)
"""

Folder Processor

"""


from .paths import (
    LOCALIZATION_FILE, LOCALIZATION_REMOTE_PATH,
    LOCALIZATION_DRIVE_PATH, LOCALIZATION_OUTPUT_PATH,
)
from . import localization_release


# 업데이트 반영
# pinisok/gaku-patcher 의 최신 release 에 첨부된 localization.json 을 기준으로
# Google Drive 의 localization.xlsx 에 새 키를 추가하고, JP 원문이 바뀐 키는 갱신,
# release 에서 사라진 키는 OBSOLETE 마커로 표시한다.
#
# 반환 시그니처는 adv.UpdateOriginalToDrive 와 동일: (file_list, warnings)
#   - file_list: [(input_path, output_path, filename)] — main.Update() 가 무시하므로
#                실제로 변경된 경우에만 한 항목을 채워 넣는다.
#   - warnings: {filename: [str, ...]} — main 의 _update_summary 가 파일별 경고로 출력.
def UpdateOriginalToDrive(bFullUpdate=False):
    release = localization_release.fetch_latest_release()
    if release is None:
        LOG_WARN(2, "Localization release fetch failed — skipping update")
        return [], {}

    cache = localization_release.load_release_cache()
    cached_tag = cache.get("tag")
    cached_sha = cache.get("asset_sha256")

    # Skip ONLY when both tag and asset sha match. A re-published release
    # with the same tag but different content (the common case for upstream
    # patch-ship hotfixes) flips the digest and triggers re-processing —
    # the previous tag-only check missed those.
    if not bFullUpdate \
            and cached_tag == release.tag \
            and cached_sha and release.asset_sha256 \
            and cached_sha == release.asset_sha256:
        LOG_INFO(2, f"Localization release {release.tag} already processed "
                    f"(same asset sha) — skip")
        return [], {}

    if cached_tag == release.tag and cached_sha != release.asset_sha256:
        LOG_INFO(2, f"Localization release {release.tag} was REPUBLISHED "
                    f"(asset sha changed) — re-processing")
    else:
        LOG_INFO(2, f"Localization release {release.tag} detected "
                    f"(previous: {cached_tag or 'none'})")

    release_json = localization_release.download_release_json(release.asset_url)
    if release_json is None:
        return [], {}

    if not os.path.exists(LOCALIZATION_DRIVE_PATH):
        LOG_WARN(2, f"{LOCALIZATION_DRIVE_PATH} not found locally — "
                    f"download from Drive (Phase 0) before Update")
        return [], {}

    diff = localization_release.diff_release_against_xlsx(
        release_json, LOCALIZATION_DRIVE_PATH
    )

    if diff.empty:
        LOG_INFO(2, f"Localization release {release.tag} has no diff vs drive xlsx — "
                    f"only the tag cache is advanced")
        localization_release.save_release_tag(release.tag, asset_sha256=release.asset_sha256)
        return [], {}

    localization_release.apply_diff_to_xlsx(release_json, diff, LOCALIZATION_DRIVE_PATH)
    localization_release.append_release_notes(release, diff)
    localization_release.save_release_tag(release.tag)

    warnings = _diff_warnings(release, diff)
    filename = os.path.basename(LOCALIZATION_DRIVE_PATH)
    file_list = [(LOCALIZATION_DRIVE_PATH, LOCALIZATION_DRIVE_PATH, filename)]
    return file_list, warnings


def _diff_warnings(release, diff) -> dict:
    """Convert a diff into per-file warnings for the gspread / log summary."""
    filename = os.path.basename(LOCALIZATION_DRIVE_PATH)
    messages = [localization_release.summarize_diff(release, diff)]
    if diff.removed:
        sample = ", ".join(diff.removed[:5])
        suffix = "" if len(diff.removed) <= 5 else f" (외 {len(diff.removed) - 5}건)"
        messages.append(f"제거된 키는 OBSOLETE 처리: {sample}{suffix}")
    if diff.changed_jp:
        sample = ", ".join(list(diff.changed_jp.keys())[:5])
        suffix = "" if len(diff.changed_jp) <= 5 else f" (외 {len(diff.changed_jp) - 5}건)"
        messages.append(f"JP 변경된 키 재검수 필요: {sample}{suffix}")
    return {filename: messages}


# 번역 수정사항 반영
# Google Drive > GakumasTranslationDataKor
def ConvertDriveToOutput(drive_file_paths=None, bFullUpdate=False):
    """Single-file convert. Triggers when explicit list non-empty OR when the
    drive xlsx's mtime is newer than the output JSON (or output missing)."""
    converted_file_list = []
    error_file_list = []

    if not os.path.exists(LOCALIZATION_DRIVE_PATH):
        LOG_INFO(2, "Localization drive xlsx not present locally, skip")
        return [], []

    explicit = bool(drive_file_paths)
    stale = False
    try:
        if not os.path.exists(LOCALIZATION_OUTPUT_PATH):
            stale = True
        elif os.path.getmtime(LOCALIZATION_DRIVE_PATH) > os.path.getmtime(LOCALIZATION_OUTPUT_PATH):
            stale = True
    except OSError:
        stale = True

    if not explicit and not stale:
        LOG_INFO(2, "Localization is up-to-date, skip")
        return [], []

    LOG_DEBUG(
        2,
        f"Localization convert — explicit={explicit} stale-by-mtime={stale}"
    )
    input_path = LOCALIZATION_DRIVE_PATH
    output_path = LOCALIZATION_OUTPUT_PATH
    LOG_DEBUG(2, f"Start convert from drive to output '{input_path}' to '{output_path}'")
    try:
        XlsxToJson(input_path, output_path)
        converted_file_list.append(os.path.basename(LOCALIZATION_DRIVE_PATH))
    except Exception as e:
        LOG_ERROR(2, f"Error during Convert localization file from drive to output: {e}")
        logger.exception(e)
        error_file_list.append((os.path.basename(LOCALIZATION_DRIVE_PATH), e))
    return error_file_list, converted_file_list