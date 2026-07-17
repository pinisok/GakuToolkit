import json
import subprocess
import sys
from pathlib import Path

import openpyxl

from scripts import masterdb2_translate


def _write_output(directory: Path, name: str, payload) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    (directory / f"{name}.json").write_text(
        json.dumps(payload, ensure_ascii=False), encoding="utf-8"
    )


def test_concat_check_fails_when_requested_file_is_missing(tmp_path):
    output = tmp_path / "output"
    output.mkdir()

    code = masterdb2_translate.concat_check(
        "DoesNotExist",
        output_json_dir=str(output),
        jp_source_dir=str(tmp_path / "jp"),
    )

    assert code < 0
    assert masterdb2_translate._concat_exit_status(code) == 2


def test_concat_check_fails_on_malformed_output_json(tmp_path):
    output = tmp_path / "output"
    output.mkdir()
    (output / "Broken.json").write_text("{not-json", encoding="utf-8")

    code = masterdb2_translate.concat_check(
        "Broken",
        output_json_dir=str(output),
        jp_source_dir=str(tmp_path / "jp"),
    )

    assert code < 0
    assert masterdb2_translate._concat_exit_status(code) == 2


def test_concat_check_passes_after_reading_a_valid_target(tmp_path):
    output = tmp_path / "output"
    _write_output(output, "Valid", {"data": [{"id": "1"}]})

    code = masterdb2_translate.concat_check(
        "Valid",
        output_json_dir=str(output),
        jp_source_dir=str(tmp_path / "jp"),
    )

    assert code == 0
    assert masterdb2_translate._concat_exit_status(code) == 0


def test_legacy_apply_command_is_disabled(tmp_path):
    translations = tmp_path / "translations.json"
    translations.write_text("{}", encoding="utf-8")
    repo_root = Path(__file__).resolve().parents[1]

    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "scripts.masterdb2_translate",
            "apply",
            str(translations),
        ],
        cwd=repo_root,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 2
    assert "apply' is disabled" in result.stderr


def _make_masterdb_xlsx(path: Path, rows: list[list]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["IMAGE", "KEY ID 0", "KEY VALUE 0", "ID", "원문", "번역", "설명"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def _draft_item(
    sheet: str,
    key_value: str,
    field_id: str,
    source: str,
    translation: str,
):
    identity = {
        "sheet": sheet,
        "keys": {"KEY ID 0": "id", "KEY VALUE 0": key_value},
        "field_id": field_id,
        "source": source,
    }
    return {
        "identity": identity,
        "identity_key": masterdb2_translate._identity_key(identity),
        "translation": translation,
    }


def test_apply_records_uses_stable_identity_for_duplicate_source_text(tmp_path):
    _make_masterdb_xlsx(
        tmp_path / "Produce.xlsx",
        [
            ["", "id", "record-a", "name", "同じ", "", ""],
            ["", "id", "record-b", "name", "同じ", "", ""],
        ],
    )
    draft = {
        "Produce": [
            _draft_item("Produce", "record-a", "name", "同じ", "첫 번째"),
            _draft_item("Produce", "record-b", "name", "同じ", "두 번째"),
        ]
    }
    draft_path = tmp_path / "draft.json"
    draft_path.write_text(json.dumps(draft, ensure_ascii=False), encoding="utf-8")

    report = masterdb2_translate.apply_record_translations(
        str(draft_path), drive_path=str(tmp_path)
    )

    assert report["failed_files"] == {}
    wb = openpyxl.load_workbook(tmp_path / "Produce.xlsx", read_only=True)
    assert wb.active["F2"].value == "첫 번째"
    assert wb.active["F3"].value == "두 번째"
    wb.close()


def test_apply_records_is_atomic_per_file_but_allows_other_files(tmp_path):
    _make_masterdb_xlsx(
        tmp_path / "Broken.xlsx",
        [["", "id", "record-a", "name", "원문", "기존", ""]],
    )
    _make_masterdb_xlsx(
        tmp_path / "Valid.xlsx",
        [["", "id", "record-b", "name", "원문", "", ""]],
    )
    draft = {
        "Broken": [
            _draft_item("Broken", "record-a", "name", "원문", "덮어쓰기")
        ],
        "Valid": [_draft_item("Valid", "record-b", "name", "원문", "새 번역")],
    }
    draft_path = tmp_path / "draft.json"
    draft_path.write_text(json.dumps(draft, ensure_ascii=False), encoding="utf-8")

    report = masterdb2_translate.apply_record_translations(
        str(draft_path), drive_path=str(tmp_path)
    )

    assert "Broken" in report["failed_files"]
    assert report["applied_files"] == {"Valid": 1}
    broken = openpyxl.load_workbook(tmp_path / "Broken.xlsx", read_only=True)
    valid = openpyxl.load_workbook(tmp_path / "Valid.xlsx", read_only=True)
    assert broken.active["F2"].value == "기존"
    assert valid.active["F2"].value == "새 번역"
    broken.close()
    valid.close()
